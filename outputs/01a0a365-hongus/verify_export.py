from pathlib import Path
import zipfile, re, xml.etree.ElementTree as ET
import openpyxl
p=Path(__file__).parent/'Hongus_Libro_de_gastos.xlsx'
ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
with zipfile.ZipFile(p) as z: parts={n:z.read(n) for n in z.namelist()}
# Remove stale imported cell payloads that the export retained in cleared inputs.
for name in list(parts):
    if re.fullmatch(r'xl/worksheets/sheet\d+\.xml',name):
        root=ET.fromstring(parts[name])
        for c in root.findall('.//s:sheetData/s:row/s:c',ns):
            m=re.fullmatch(r'([A-Z]+)(\d+)',c.attrib['r'])
            if m[1] in ['B','C','D','E','F','G','H','J','K','L','M','N'] and 16<=int(m[2])<=115:
                for child in list(c):
                    if child.tag.split('}')[-1] in ['f','v','is']: c.remove(child)
                c.attrib.pop('t',None)
        parts[name]=ET.tostring(root,encoding='utf-8',xml_declaration=True)
with zipfile.ZipFile(p,'w',zipfile.ZIP_DEFLATED) as z:
    for n,b in parts.items(): z.writestr(n,b)
w=openpyxl.load_workbook(p)
s=w.active
assert all(s.cell(r,c).value is None for r in range(16,116) for c in [2,3,4,5,6,7,8,10,11,12,13,14])
assert s['I16'].data_type=='f' and s['O115'].data_type=='f'
assert len(s.data_validations.dataValidation)==5
assert len(s.tables)==1
print('Verified: blank inputs, formulas, validations, table and workbook reopening.')
