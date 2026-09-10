from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
OUT=Path(__file__).parent
ASSETS=ROOT/'Assets/Hongus/Identidad-visual-v1'
def save(path,s): (ROOT/path).write_text(s.strip()+'\n',encoding='utf-8')

# Retain current decisions without historical references.
for folder in [p for p in ROOT.iterdir() if p.is_dir() and p.name[:2].isdigit()]:
 for path in folder.rglob('*.md'):
  text=path.read_text(encoding='utf-8')
  text=text.replace('[[Propuesta original y evolucion]]','[[Identidad de Hongus]]')
  lines=[l for l in text.splitlines() if not any(x in l for x in ['Fuente histórica conservada','El Word propone 10,000','El Word promete aumentar','Las cifras históricas anteriores','Versiones anteriores en 99-Archive.'])]
  path.write_text('\n'.join(lines)+'\n',encoding='utf-8')
save('00-Inbox/Brief fundador.md','''# Brief fundador
Información vigente al 10 de septiembre de 2026.

Hongus es una plataforma de empleabilidad verde en preparación para todo México y personas de 18 años en adelante. Su público principal son estudiantes de preparatoria y superior, egresados y profesionales sin experiencia. Profesionales experimentados participarán como mentores y guías.

Nombre y lema aprobados: **Hongus — Construye experiencia. Cultiva futuro.** Identidad gráfica aprobada en [[Identidad visual de Hongus]].

Planes aceptados en [[Planes y beneficios]]: Gratis, Estudiante $50 MXN/mes e Inicio Profesional $200 MXN/mes. Repartos para MICE-LO en [[Modelo de negocio]]. Reglas de [[Mentoria one to one]] y [[Contactos y Hongus Verify]].

No hay organizaciones ni mentores confirmados. Primero identidad; después arquitectura, equipos, desarrollo, tickets y metodología. Stack, presupuesto, fechas y alcance de lanzamiento pendientes.
''')
save('03-Architecture/Decisiones.md','''# Decisiones vigentes
Actualización al 10 de septiembre de 2026. Fuente: decisiones del fundador en conversación.

| Área | Decisión vigente |
| --- | --- |
| Marca | Hongus y lema Construye experiencia. Cultiva futuro. aprobados. |
| Visual | Logotipo, paleta, tipografía y aplicaciones de [[Identidad visual de Hongus]] aprobados. |
| Público | México, desde 18 años; estudiantes, egresados y profesionales sin experiencia. Mentores experimentados como acompañantes. |
| Planes | [[Planes y beneficios]]: $50 y $200 MXN mensuales; 1 y 4 mentorías. Gratis: 3 postulaciones/mes. |
| Aportaciones | [[Modelo de negocio]]: 20% de ingresos propios a MICE-LO; extras con reparto exclusivo 75% mentor y 25% MICE-LO. |
| Confianza | [[Contactos y Hongus Verify]]: canal de asuntos y logros confirmados por responsables. |
| Organización del trabajo | Identidad, arquitectura, equipos, desarrollo, tickets y metodología. |

La arquitectura, stack y alcance técnico del lanzamiento no están aprobados. La base de conocimiento conserva la versión vigente a solicitud del fundador.
''')
save('02-Product/Users/Talento.md','''# Talento
Público vigente: personas de 18 años en adelante en México.

- Estudiantes de preparatoria y educación superior: construir experiencia durante sus estudios; condición verificada.
- Egresados y profesionales sin experiencia: preparar su incorporación laboral, mostrar logros y recibir acompañamiento.
- Profesionales con experiencia: participan principalmente como mentores y guías, no como público prioritario del servicio de primera experiencia.

Ver [[Planes y beneficios]], [[Identidad de Hongus]] y [[Mentoria one to one]]. Las funciones técnicas específicas se definirán durante la arquitectura.
''')
save('02-Product/Users/Tipos de Usuario.md','''# Tipos de Usuario
| Participante | Papel |
| --- | --- |
| [[Talento]] | Estudiantes mayores de edad, egresados y profesionales sin experiencia. |
| [[Empresa]] | Ofrecer oportunidades y confirmar contratación o logros. |
| [[Institucion Educativa]] | Ofrecer proyectos, prácticas y oportunidades; confirmar resultados. |
| [[Mentor]] | Acompañamiento individual y confirmación de actividades terminadas. |
| [[Organizacion]] | Empresas, universidades e instituciones participantes. |
| Administración de Hongus | Atención de errores y reclamaciones de Hongus Verify. |

Gobierno fuera del público inicial. Alianzas aún sin confirmar. Roles y permisos técnicos por definir. Ver [[Contactos y Hongus Verify]].
''')
p=ROOT/'Home.md';s=p.read_text(encoding='utf-8').replace('[[Vision General]] y [[Propuesta original y evolucion]]','[[Vision General]] y [[Identidad de Hongus]]').replace('Definir [[Identidad de Hongus]] y elegir un recorrido en [[MVP]].','Consultar [[Identidad de Hongus]] e [[Identidad visual de Hongus]] aprobada; después definir el recorrido de lanzamiento.')
s += '\n## Presentación de Hongus\n[[entregables/hongus-identidad/Hongus_Presentacion_Institucional.docx|Descargar presentación completa en Word]]\n'
p.write_text(s,encoding='utf-8')
p=ROOT/'01-Vision/Objetivos.md';s=p.read_text(encoding='utf-8').replace('profesionales con y sin experiencia','profesionales sin experiencia');s+='\nAspiración del primer año: promover conciencia ambiental en al menos 10,000 usuarios y vincular organizaciones. Método de medición y meta de contrataciones pendientes.\n';p.write_text(s,encoding='utf-8')
p=ROOT/'01-Vision/Propuesta de Valor.md';s=p.read_text(encoding='utf-8').replace('Empresas públicas y privadas','Empresas participantes');p.write_text(s,encoding='utf-8')
p=ROOT/'02-Product/MVP.md';s=p.read_text(encoding='utf-8');s=s.split('## Público confirmado')[0]+'## Público\nVer [[Talento]] para el público vigente. El alcance técnico de lanzamiento se definirá después de la identidad.\n';s=s.replace('Una empresa pública o privada','Una organización participante');p.write_text(s,encoding='utf-8')
p=ROOT/'05-Design-UX/Identidad visual de Hongus.md';s=p.read_text(encoding='utf-8').replace('La lámina conserva su rótulo de propuesta como versión presentada y aprobada.','La lámina muestra la identidad aprobada.');p.write_text(s,encoding='utf-8')
p=ASSETS/'hongus-lamina-visual.svg';s=p.read_text(encoding='utf-8').replace('PROPUESTA V1','IDENTIDAD APROBADA').replace('Propuesta visual pendiente de aprobación','Identidad visual aprobada');p.write_text(s,encoding='utf-8')

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.8)
for name,size in [('Normal',11),('Title',30),('Subtitle',13),('Heading 1',21),('Heading 2',13)]:
 st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(8)
doc.styles['Normal'].paragraph_format.line_spacing=1.12
doc.styles['Heading 2'].paragraph_format.space_before=Pt(12)
footer=sec.footer.paragraphs[0];footer.alignment=2
footer.add_run('Hongus  |  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
for r in footer.runs:r.font.size=Pt(9)
doc.core_properties.title='Hongus Presentación institucional'
doc.core_properties.author='Hongus'
doc.core_properties.subject='Identidad, propuesta de valor, servicios y colaboración'
def p(s):doc.add_paragraph(s)
def h(s):doc.add_heading(s,2)
def page(s):doc.add_page_break();doc.add_heading(s,1)
def bullet(s):doc.add_paragraph(s,'List Bullet')
def table(headers,rows,widths):
 t=doc.add_table(rows=1, cols=len(headers));t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for c,v in zip(t.rows[0].cells,headers):c.text=v
 for row in rows:
  cells=t.add_row().cells
  for c,v in zip(cells,row):c.text=v
 for i,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if i==0:pr.append(OxmlElement('w:tblHeader'))
  for j,c in enumerate(row.cells):
   c.width=Inches(widths[j]);c.vertical_alignment=1
   tc=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'303C38' if i==0 else ('F2F5F3' if i%2==0 else 'FFFFFF'));tc.append(sh)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   tc.append(borders);m=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'100');e.set(qn('w:type'),'dxa');m.append(e)
   tc.append(m)
   for q in c.paragraphs:
    q.paragraph_format.space_after=Pt(3);q.paragraph_format.line_spacing=1.05
    for r in q.runs:r.font.size=Pt(10);r.bold=i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '102A22')
 return t

doc.add_picture(str(ASSETS/'hongus-logo.png'),width=Inches(3.7))
doc.add_paragraph('Presentación de Hongus','Title')
doc.add_paragraph('Identidad y propuesta para nuestra comunidad y colaboradores','Subtitle')
p('Construye experiencia. Cultiva futuro.')
h('Una oportunidad para comenzar')
p('Hongus es una plataforma de empleabilidad verde en preparación para México. Nace para que estudiantes mayores de edad, egresados y profesionales sin experiencia puedan desarrollar sus capacidades, demostrar sus logros y acercarse a oportunidades laborales con conciencia ambiental.')
p('Reuniremos oportunidades, proyectos, mentorías individuales y herramientas para construir un CV con ayuda de inteligencia artificial. Empresas, universidades, instituciones y profesionales podrán participar ofreciendo espacios de aprendizaje, acompañamiento y confirmación de resultados.')
h('Qué encontrará en este documento')
p('Nuestra razón de ser, las personas a quienes buscamos ayudar, los servicios previstos, las suscripciones, Hongus Verify, el compromiso con MICE-LO y la identidad visual aprobada. También presentamos las formas en que futuros colaboradores podrán sumarse.')
h('Dónde estamos')
p('Estamos en la etapa de definición de identidad, previa a la arquitectura, conformación de equipos y desarrollo. El alcance previsto es nacional y el acceso inicial será para personas de 18 años en adelante. Aún no contamos con organizaciones o mentores confirmados.')
p('México · Septiembre de 2026')

page('Propósito y personalidad')
h('Por qué nace Hongus')
p('Hongus surge de una experiencia concreta: enfrentar procesos de contratación difíciles por no contar con experiencia y encontrar pocas oportunidades para adquirirla. Creemos que los estudiantes pueden construir proyectos profesionales cuando cuentan con oportunidades y acompañamiento.')
h('Propósito')
p('Hacer que la falta de experiencia deje de cerrar puertas y que el crecimiento profesional contribuya al cuidado del planeta.')
h('Misión')
p('Acompañar al talento emergente en la construcción de experiencia profesional mediante oportunidades, mentoría y logros verificables, promoviendo la conciencia ambiental en cada etapa.')
h('Visión')
p('Ser una plataforma líder en empleabilidad verde, reconocida por abrir oportunidades al talento emergente y contribuir al cuidado ambiental junto con su comunidad y MICE-LO.')
h('Lo que nos guía')
for s in ['Oportunidad: abrir espacios para demostrar el potencial de quienes comienzan.','Profesionalismo y honestidad: cumplir compromisos y comunicar con claridad alcances y resultados.','Colaboración y respeto: aprender con apoyo, compartir conocimiento y tratar a cada persona con dignidad.','Compromiso ambiental: vincular el desarrollo profesional con acciones de cuidado del entorno.']:bullet(s)
h('Cómo queremos hacerte sentir')
p('Con confianza, alegría, entusiasmo y pertenencia. Hongus tendrá una voz cercana, juvenil e incluyente, libre de expresiones racistas o elitistas. Evitaremos una comunicación solemne o distante y responderemos con responsabilidad cuando algo salga mal.')

page('Comunidad y oportunidades')
h('Para quién construimos Hongus')
table(['Participante','Cómo participará'],[
('Estudiantes de 18 años o más','Preparatoria y educación superior. Construcción de experiencia compatible con su formación.'),
('Egresados y profesionales sin experiencia','Preparación de candidaturas, desarrollo de capacidades y búsqueda del primer empleo.'),
('Profesionales con experiencia','Mentoría y guía para quienes comienzan.'),
('Empresas, universidades e instituciones','Oportunidades, proyectos, prácticas y confirmación de los logros correspondientes.')],[2.2,4.4])
h('Qué entendemos por oportunidades verdes')
p('Hongus contempla actividades con objetivos ambientales directos, trabajos que apoyen a organizaciones ambientales —como desarrollar una herramienta digital— y puestos tradicionales dentro de organizaciones con iniciativas sostenibles. Cada oportunidad deberá explicar su relación con el cuidado ambiental.')
p('El catálogo previsto incluye empleos, prácticas profesionales, servicio social, convocatorias y proyectos. La admisión dependerá de los requisitos y cupos de cada organización.')
h('Del aprendizaje a la experiencia')
for s in ['Descubrir una oportunidad y comprender sus requisitos.','Participar en un proyecto o actividad y recibir acompañamiento cuando se necesite.','Completar el trabajo y obtener confirmación de los logros.','Incorporar la experiencia al CV con ayuda de IA.','Postularse a un empleo y registrar la contratación cuando la organización la confirme.']:doc.add_paragraph(s,'List Number')
p('Cada persona podrá avanzar según su preparación; cursar todo el recorrido no será una condición para postularse. La plataforma facilitará oportunidades y acompañamiento, sin garantizar contratación.')

page('Suscripciones para cada etapa')
p('Los planes individuales se han definido en pesos mexicanos y con periodicidad mensual. El plan Estudiante requiere una constancia de estudios o credencial vigente, con revisión periódica.')
table(['Beneficio','Gratis','Estudiante','Inicio Profesional'],[
('Precio mensual','$0','$50 MXN','$200 MXN'),
('Postulaciones a empleos','3 al mes','Ilimitadas','Ilimitadas'),
('Convocatorias y proyectos','No incluidos','Acceso ilimitado','Acceso ilimitado'),
('Perfil y contactos','No incluidos','Formación, intereses y proyectos','Capacidades, portafolio y disponibilidad'),
('CV con IA','No incluido','Un CV principal editable','CV principal y versiones por vacante'),
('Mentorías individuales','No incluidas','1 al mes','4 al mes'),
('Guías y recursos','No incluidos','Prácticas, servicio social, portafolio y conciencia ambiental','Entrevistas, postulaciones y presentación de experiencia'),
('Propuestas y consultas','No incluidas','Incluidas','Incluidas'),
('Hongus Verify','No incluido','Logros completados y confirmados','Logros completados y confirmados'),
('Mentorías adicionales','No incluidas','Con costo extra','Con costo extra')],[1.55,.85,2.1,2.1])
h('Acceso con condiciones claras')
p('El acceso ilimitado permite consultar y postularse sin un límite de uso establecido por el plan. No sustituye los requisitos de cada oportunidad. Los documentos de Hongus Verify requieren completar y confirmar la actividad; no se obtienen solo por pagar.')
p('La duración de las mentorías la fija cada mentor. Las sesiones incluidas vencen al terminar el mes de suscripción y no se acumulan. Las condiciones finales de cobro, impuestos y cancelación se precisarán antes del lanzamiento.')

page('Acompañamiento y confianza')
h('Mentorías individuales')
p('Las sesiones permitirán trabajar objetivos concretos: preparar una candidatura, presentar proyectos o fortalecer el perfil. Los mentores participarán voluntariamente en las sesiones incluidas en las suscripciones. La duración y disponibilidad se mostrarán antes de reservar.')
p('Para sesiones adicionales, el mentor establecerá el precio. Recibirá el 75% del importe y el 25% se destinará a MICE-LO para reforestación. Este reparto es exclusivo y no lleva una aportación adicional del 20%.')
h('Contactos con un propósito')
p('Entre personas, Hongus contempla solicitudes de contacto que cada usuario podrá aceptar o rechazar. Con empresas, universidades e instituciones existirá un canal de propuestas y consultas: un asunto, un motivo y una descripción, con archivos cuando correspondan.')
p('Ambas partes podrán responder dentro del mismo asunto para presentar ideas, debatir temas y proponer colaboraciones. La interacción estará enfocada en oportunidades y acompañamiento.')
h('Hongus Verify')
p('Cuando una persona culmine una convocatoria, proyecto, mentoría o guía, el responsable confirmará el logro. En el caso del primer empleo, la organización contratante confirmará la contratación. El registro incluirá actividad, fechas, resultado o entregable y nombre del responsable.')
p('Con esa confirmación, Hongus Verify emitirá el documento correspondiente e incorporará los logros al recorrido profesional del usuario. La IA apoyará el procesamiento de la información; la confirmación recaerá en la persona u organización responsable.')
p('Los documentos podrán adoptar la forma de constancia o reconocimiento y, para el primer empleo, se contempla un certificado cuyo alcance está por definir. El reconocimiento ante la SEP se encuentra en gestión y no se presenta como obtenido.')
h('Atención de errores')
p('La revisión de errores y reclamaciones corresponderá a una persona del equipo interno con perfil de administración de empresas, inicialmente practicante y posteriormente profesional. Los procedimientos de revisión, corrección y revocación se desarrollarán antes de operar el servicio.')

page('Compromiso ambiental y colaboración')
h('Hongus y MICE LO')
p('Hongus y MICE-LO se conciben como organizaciones hermanas y separadas. El modelo de Hongus contempla destinar el 20% de sus ingresos propios, incluidas las suscripciones, a MICE-LO. El 80% restante se destinará a reinversión y pago de nómina.')
table(['Origen','Destino definido'],[('Ingresos propios de Hongus','20% MICE-LO / 80% reinversión y nómina'),('Mentorías adicionales','25% MICE-LO para reforestación / 75% mentor')],[2.4,4.2])
p('La base de cálculo, las transferencias y la forma de reportar las aportaciones se precisarán en la operación. Las actividades y aportaciones realizadas podrán difundirse en las páginas oficiales de Hongus y MICE-LO con información que las respalde.')
h('Qué buscamos construir con colaboradores')
for s in ['Empresas: oportunidades de entrada y proyectos con requisitos claros, acompañamiento y confirmación de resultados.','Universidades e instituciones: vinculación con estudiantes mayores de edad y espacios de prácticas, servicio social o proyectos.','Mentores y guías: conocimientos, tiempo y orientación para quienes necesitan dar sus primeros pasos.','Colaboradores de Hongus: capacidades para definir y construir la plataforma y su operación. Los equipos se organizarán en la siguiente etapa.']:bullet(s)
h('Nuestra aspiración')
p('Queremos promover conciencia ambiental en al menos 10,000 usuarios durante el primer año y sumar empresas, universidades e instituciones que abran oportunidades. La metodología de medición y la meta de contrataciones están por definir. Estas cifras representan una aspiración, no resultados alcanzados.')
h('Próxima etapa')
p('Con la identidad definida, continuaremos con arquitectura, equipos, desarrollo, sistema de tickets y metodología de trabajo. La incorporación de aliados y mentores y la validación de costos y capacidad serán necesarias para preparar el lanzamiento.')

page('Identidad visual de Hongus')
p('El nombre combina hongo y us, “nosotros”. El símbolo integra un hongo tecnológico a la inicial h; la palabra se completa con ongus. Su forma orgánica expresa crecimiento compartido y el pequeño circuito aporta el carácter tecnológico.')
doc.add_picture(str(ASSETS/'hongus-lamina-visual.png'),width=Inches(6.55))

page('Expresión y uso de la marca')
h('Paleta aprobada')
table(['Color','Código','Aplicación'],[('Bosque','#145C43','Marca y acciones principales'),('Tinta','#102A22','Texto y fondos oscuros'),('Lima','#D5F67A','Acentos y fondos destacados'),('Niebla','#F5F8F2','Fondo general'),('Blanco','#FFFFFF','Superficies y marca inversa')],[1.35,1.35,3.9])
h('Tipografía y composición')
p('Arial Bold para titulares y marca; Arial Regular para lectura e interfaz. Formas simples, espacios amplios y esquinas suaves acompañan una identidad minimalista, profesional, juvenil y ambiental.')
h('Uso del logotipo')
p('Conservar proporciones y un espacio libre mínimo equivalente a un cuarto de la altura del símbolo. Usar la variante principal sobre fondos claros y la inversa sobre fondos oscuros. Evitar estirarlo, rotarlo, añadir sombras o colocarlo sobre fondos de poco contraste.')
p('Como referencia digital inicial: 160 px de ancho mínimo para el conjunto y 32 px para el símbolo, revisando la legibilidad del circuito a pequeña escala. Antes de imprenta, preparar el archivo con letras convertidas a contornos.')
h('Nuestra voz')
p('Hablamos de tú con alegría y respeto. La inclusión se expresa en cómo damos la bienvenida, explicamos requisitos y acompañamos a cada persona. Usamos mensajes concretos y evitamos promesas de empleo o acreditaciones que no podamos respaldar.')
for s in ['“Tu talento tiene un lugar aquí.”','“Da tu siguiente paso con apoyo.”','“Cada logro cuenta. Cada acción por el planeta también.”']:bullet(s)
h('Lema oficial')
p('Construye experiencia. Cultiva futuro.')
p('Hongus invita a estudiantes, profesionales y organizaciones a contribuir a una comunidad donde aprender, demostrar capacidades y cuidar el entorno formen parte del mismo camino.')
doc.save(OUT/'Hongus_Presentacion_Institucional.docx')
print('Documento creado:',OUT/'Hongus_Presentacion_Institucional.docx')
